import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def add_inline_runs(paragraph, text, default_color=RGBColor(30, 41, 59), font_size=Pt(10.5), font_name="Calibri"):
    pattern = re.compile(r'(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\)|[^\`\*\[]+|\[|\]|\*|\`)')
    tokens = pattern.findall(text)
    
    for token in tokens:
        if not token:
            continue
        if token.startswith('`') and token.endswith('`') and len(token) >= 2:
            code_text = token[1:-1]
            run = paragraph.add_run(code_text)
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(225, 29, 72)
            rPr = run._r.get_or_add_rPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
            rPr.append(shd)
        elif token.startswith('**') and token.endswith('**') and len(token) >= 4:
            bold_text = token[2:-2]
            run = paragraph.add_run(bold_text)
            run.font.name = font_name
            run.font.size = font_size
            run.font.bold = True
            run.font.color.rgb = default_color
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            italic_text = token[1:-1]
            run = paragraph.add_run(italic_text)
            run.font.name = font_name
            run.font.size = font_size
            run.font.italic = True
            run.font.color.rgb = default_color
        elif token.startswith('[') and '](' in token and token.endswith(')'):
            m = re.match(r'\[(.*?)\]\((.*?)\)', token)
            if m:
                link_text, link_url = m.groups()
                run = paragraph.add_run(link_text)
                run.font.name = font_name
                run.font.size = font_size
                run.font.underline = True
                run.font.color.rgb = RGBColor(26, 86, 219)
            else:
                run = paragraph.add_run(token)
                run.font.name = font_name
                run.font.size = font_size
                run.font.color.rgb = default_color
        else:
            run = paragraph.add_run(token)
            run.font.name = font_name
            run.font.size = font_size
            run.font.color.rgb = default_color

def create_specification_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("UniProcure — Frontend System Specification | Confidential & Institutional")
        f_run.font.name = "Calibri"
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = RGBColor(148, 163, 184)

    NAVY_PRIMARY = RGBColor(26, 86, 219)    # #1A56DB
    SLATE_DARK = RGBColor(15, 23, 42)      # #0F172A
    SLATE_TEXT = RGBColor(51, 65, 85)      # #334155

    i = 0
    n = len(lines)
    in_code_block = False
    code_block_lines = []
    code_block_lang = ""

    while i < n:
        raw_line = lines[i]
        line = raw_line.rstrip('\r\n')

        if line.strip().startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_block_lang = line.strip()[3:].strip()
                code_block_lines = []
                i += 1
                continue
            else:
                in_code_block = False
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "F8FAFC")
                set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
                
                tcPr = cell._tc.get_or_add_tcPr()
                tcBorders = parse_xml(
                    f'<w:tcBorders {nsdecls("w")}>'
                    f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                    f'<w:left w:val="single" w:sz="18" w:space="0" w:color="1A56DB"/>'
                    f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                    f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                    f'</w:tcBorders>'
                )
                tcPr.append(tcBorders)

                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(2)
                cp.paragraph_format.space_after = Pt(2)
                cp.paragraph_format.line_spacing = 1.15

                code_text = "\n".join(code_block_lines)
                c_run = cp.add_run(code_text)
                c_run.font.name = "Consolas"
                c_run.font.size = Pt(8.5)
                c_run.font.color.rgb = RGBColor(30, 41, 59)

                sp_p = doc.add_paragraph()
                sp_p.paragraph_format.space_before = Pt(0)
                sp_p.paragraph_format.space_after = Pt(4)

                i += 1
                continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        if line.strip().startswith('|') and '|' in line.strip()[1:]:
            table_lines = []
            while i < n and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1

            if len(table_lines) >= 2:
                header_raw = [c.strip() for c in table_lines[0].strip('|').split('|')]
                data_rows = []
                for t_line in table_lines[2:]:
                    row = [c.strip() for c in t_line.strip('|').split('|')]
                    if len(row) < len(header_raw):
                        row += [""] * (len(header_raw) - len(row))
                    data_rows.append(row[:len(header_raw)])

                tbl = doc.add_table(rows=len(data_rows) + 1, cols=len(header_raw))
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(tbl, color="CBD5E1", sz="6", val="single")

                # Column widths heuristic
                if len(header_raw) == 4:
                    col_widths = [Inches(1.5), Inches(1.8), Inches(1.2), Inches(2.2)]
                else:
                    col_widths = [Inches(6.7 / len(header_raw))] * len(header_raw)

                hdr_cells = tbl.rows[0].cells
                for col_idx, col_name in enumerate(header_raw):
                    cell = hdr_cells[col_idx]
                    cell.width = col_widths[col_idx]
                    set_cell_background(cell, "1A56DB")
                    set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    add_inline_runs(p, col_name, default_color=RGBColor(255, 255, 255), font_size=Pt(9.5))
                    for r in p.runs:
                        r.font.bold = True

                for row_idx, r_data in enumerate(data_rows):
                    row_cells = tbl.rows[row_idx + 1].cells
                    fill_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                    for col_idx, cell_value in enumerate(r_data):
                        cell = row_cells[col_idx]
                        cell.width = col_widths[col_idx]
                        set_cell_background(cell, fill_color)
                        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.line_spacing = 1.15
                        add_inline_runs(p, cell_value, default_color=RGBColor(30, 41, 59), font_size=Pt(9))

                sp_p = doc.add_paragraph()
                sp_p.paragraph_format.space_before = Pt(0)
                sp_p.paragraph_format.space_after = Pt(6)
            continue

        if not line.strip():
            i += 1
            continue

        if line.strip() in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="E2E8F0"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            i += 1
            continue

        if line.startswith('# '):
            h_text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = "Calibri"
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = NAVY_PRIMARY
            i += 1
            continue

        if line.startswith('## '):
            h_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = "Calibri"
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = SLATE_DARK
            i += 1
            continue

        if line.startswith('### '):
            h_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = "Calibri"
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = NAVY_PRIMARY
            i += 1
            continue

        if line.startswith('#### '):
            h_text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = SLATE_DARK
            i += 1
            continue

        list_match = re.match(r'^(\s*)([\*\-])\s+(.*)$', line)
        if list_match:
            indent_spaces = len(list_match.group(1))
            bullet_text = list_match.group(3)
            
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            if indent_spaces >= 2:
                p.paragraph_format.left_indent = Inches(0.4)
                b_run = p.add_run("◦  ")
                b_run.font.color.rgb = NAVY_PRIMARY
            else:
                p.paragraph_format.left_indent = Inches(0.2)
                b_run = p.add_run("▪  ")
                b_run.font.color.rgb = NAVY_PRIMARY
                b_run.font.size = Pt(9)
            
            add_inline_runs(p, bullet_text, default_color=SLATE_TEXT, font_size=Pt(10))
            i += 1
            continue

        num_list_match = re.match(r'^(\s*)(\d+)\.\s+(.*)$', line)
        if num_list_match:
            indent_spaces = len(num_list_match.group(1))
            num_str = num_list_match.group(2)
            item_text = num_list_match.group(3)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Inches(0.25 if indent_spaces < 2 else 0.45)
            
            n_run = p.add_run(f"{num_str}. ")
            n_run.font.bold = True
            n_run.font.color.rgb = NAVY_PRIMARY
            add_inline_runs(p, item_text, default_color=SLATE_TEXT, font_size=Pt(10))
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_inline_runs(p, line, default_color=SLATE_TEXT, font_size=Pt(10.5))
        i += 1

    doc.save(docx_path)
    print(f"Successfully generated docx: {docx_path}")

if __name__ == "__main__":
    src = "/Users/radwanabirgmail.com/Downloads/uni_procurement/FRONTEND_SYSTEM_SPECIFICATION.md"
    dst = "/Users/radwanabirgmail.com/Downloads/uni_procurement/FRONTEND_SYSTEM_SPECIFICATION.docx"
    create_specification_docx(src, dst)
